"""Read the local BiGG inventory workbook with stdlib fallback support."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree

_NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _text(element: ElementTree.Element | None) -> str:
    return "" if element is None else "".join(element.itertext())


def _read_xml_rows(path: Path, sheet_name: str) -> list[dict[str, object]]:
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        shared: list[str] = []
        if "xl/sharedStrings.xml" in names:
            shared_root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            shared = [_text(item) for item in shared_root.findall(f"{{{_NS_MAIN}}}si")]
        workbook_root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        rels_root = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {
            relation.attrib["Id"]: relation.attrib["Target"]
            for relation in rels_root
            if relation.attrib.get("Type", "").endswith("/worksheet")
        }
        target = None
        for sheet in workbook_root.findall(f".//{{{_NS_MAIN}}}sheet"):
            if sheet.attrib.get("name") == sheet_name:
                target = targets.get(sheet.attrib.get(f"{{{_NS_REL}}}id"))
                break
        if target is None:
            raise ValueError(f"workbook sheet is missing: {sheet_name}")
        sheet_path = target.lstrip("/")
        if not sheet_path.startswith("xl/"):
            sheet_path = f"xl/{sheet_path}"
        root = ElementTree.fromstring(archive.read(sheet_path))
        rows: list[dict[str, object]] = []
        headers: list[str] = []
        for row in root.findall(f".//{{{_NS_MAIN}}}row"):
            values: dict[int, str] = {}
            for cell in row.findall(f"{{{_NS_MAIN}}}c"):
                reference = cell.attrib.get("r", "A1")
                match = re.match(r"[A-Z]+", reference, re.I)
                if match is None:
                    continue
                column = 0
                for character in match.group(0):
                    column = column * 26 + ord(character.upper()) - ord("A") + 1
                cell_type = cell.attrib.get("t")
                if cell_type == "inlineStr":
                    value = _text(cell.find(f"{{{_NS_MAIN}}}is"))
                else:
                    value = _text(cell.find(f"{{{_NS_MAIN}}}v"))
                    if cell_type == "s" and value:
                        value = shared[int(value)]
                values[column] = value
            if not headers:
                headers = [values.get(index, "") for index in range(1, max(values, default=0) + 1)]
                continue
            rows.append(
                {
                    header: values.get(index, "")
                    for index, header in enumerate(headers, 1)
                    if header
                }
            )
    return rows


def read_inventory_rows(path: Path, sheet_name: str = "原始数据") -> list[dict[str, object]]:
    """Read rows using openpyxl when installed, otherwise stdlib XML."""
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError:
        return _read_xml_rows(path, sheet_name)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        headers = [str(value) for value in next(sheet.iter_rows(values_only=True))]
        return [
            dict(zip(headers, values, strict=True))
            for values in sheet.iter_rows(min_row=2, values_only=True)
        ]
    finally:
        workbook.close()
